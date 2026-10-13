#!/usr/bin/env python3.13
"""Готовит пакет на рамочное соглашение по металлу, Рок-Айленд, W519TC26BOAS.

Делает два файла с именами, которые требует раздел L солиситации:

    ISP GROUP LLC_Att001_Technical.xlsx   опросник по номенклатуре
    ISP GROUP LLC_Att002_POCs.docx        контактные лица

Оба собираются из подлинных шаблонов заказчика, скачанных с SAM, чтобы
формат остался их. В колонтитул каждой страницы ставится имя участника,
CAGE, номер солиситации и дата, как требует раздел L, пункт 4.

    python3.13 tools/boa-rock-island.py <папка-с-вложениями> [куда]

Вложения качаются так:

    curl -H 'Accept: application/hal+json' \\
      'https://sam.gov/api/prod/opps/v3/opportunities/3bb8c255f2384fc0abbd8cbbf79513c4/resources'
"""

import datetime
import shutil
import sys
from pathlib import Path

import docx
import openpyxl

NAME = "ISP GROUP LLC"
UEI = "WA46PAMTYMQ7"
CAGE = "217F4"
ADDRESS = "16395 Biscayne Blvd Apt 818, Aventura, FL 33160-5743, United States"
CONTACT = "Aleksandr Orlov, Member"
EMAIL = "info@ispgroupgc.com"
PHONE = "(929) 830-0038"
RFP = "W519TC26BOAS"

# Позиции, по которым отвечаем НЕТ. Номера строк листа Sheet1.
# 413: стальная дробь, требуется сертификат конкретного изготовителя
# (Metaltec). Пока у нас нет подтверждённого канала к нему, честный ответ НЕТ.
NO_ROWS = {413}

SRC_XLSX = "Attachment 001 Rev 1 PURCHASE DESCRIPTION.xlsx"
SRC_DOCX = "Attachment 002 Offeror Points of Contact (POCs) Template.docx"


def stamp() -> str:
    today = datetime.date.today().strftime("%d %B %Y")
    return f"{NAME}  |  CAGE {CAGE}  |  {RFP}  |  {today}"


def build_xlsx(src: Path, out: Path) -> tuple[int, int]:
    wb = openpyxl.load_workbook(src)
    ws = wb["Sheet1"]

    ws["B1"] = NAME
    ws["B2"] = UEI
    ws["B3"] = CAGE

    yes = no = 0
    for row in range(5, ws.max_row + 1):
        item = ws.cell(row=row, column=1).value
        if item is None or not str(item).strip():
            continue
        if row in NO_ROWS:
            ws.cell(row=row, column=3).value = "X"
            no += 1
        else:
            ws.cell(row=row, column=2).value = "X"
            yes += 1

    ws.oddHeader.left.text = stamp()
    ws.oddFooter.right.text = "&P / &N"
    wb.save(out)
    return yes, no


def build_docx(src: Path, out: Path) -> None:
    shutil.copy(src, out)
    d = docx.Document(str(out))

    values = {
        "Company Name:": NAME,
        "UEI:": UEI,
        "CAGE:": CAGE,
        "Address:": ADDRESS,
        "Name:": CONTACT,
        "Email:": EMAIL,
        "Phone Number:": PHONE,
    }
    for p in d.paragraphs:
        label = p.text.strip()
        for key, val in values.items():
            if label == key.rstrip() or label == key:
                # дописываем в последний прогон, чтобы не потерять оформление
                p.runs[-1].text = p.runs[-1].text.rstrip() + f" {val}"
                break

    for section in d.sections:
        head = section.header.paragraphs[0]
        head.text = stamp()

    d.save(out)


def main() -> None:
    src = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
    out = Path(sys.argv[2] if len(sys.argv) > 2 else "bids/W519TC26BOAS")
    out.mkdir(parents=True, exist_ok=True)

    for f in (SRC_XLSX, SRC_DOCX):
        if not (src / f).exists():
            raise SystemExit(f"Нет исходного шаблона: {src / f}")

    yes, no = build_xlsx(src / SRC_XLSX, out / f"{NAME}_Att001_Technical.xlsx")
    build_docx(src / SRC_DOCX, out / f"{NAME}_Att002_POCs.docx")

    total = yes + no
    print(f"позиций в опроснике: {total}")
    print(f"  ДА:  {yes}  ({yes / total:.0%})")
    print(f"  НЕТ: {no}")
    print(f"порог допуска: 50 процентов, {'пройден' if yes / total >= 0.5 else 'НЕ ПРОЙДЕН'}")
    print(f"записано в {out}")


if __name__ == "__main__":
    main()
